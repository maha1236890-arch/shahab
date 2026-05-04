"""
app.py – واجهة ويب لنظام الحضور والغياب
يعمل على أي جهاز (هاتف / حاسوب / لوحي) عبر المتصفح
"""

import os
import sys
import tempfile
import functools
import threading
from collections import Counter
from datetime import date, datetime

from flask import (
    Flask, render_template, request, redirect,
    url_for, session, flash, send_file, jsonify,
)

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import database as db
import export_utils as eu

# ── تهيئة التطبيق ──────────────────────────────────────────────────────────
app = Flask(__name__)

@app.template_filter('time12')
def time12_filter(value):
    """تحويل HH:MM إلى تنسيق 12 ساعة: 8:30 ص / 2:45 م"""
    if not value or str(value).strip() in ('', '—', 'None'):
        return '—'
    try:
        t = datetime.strptime(str(value).strip(), '%H:%M')
        h, m = t.hour, t.minute
        period = 'م' if h >= 12 else 'ص'
        h12 = h % 12 or 12
        return f'{h12}:{m:02d} {period}'
    except Exception:
        return str(value)

app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    PERMANENT_SESSION_LIFETIME=86400 * 7,   # أسبوع
    TEMPLATES_AUTO_RELOAD=True,
)

_KEY_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".secret_key")
if os.path.exists(_KEY_FILE):
    with open(_KEY_FILE, "rb") as _f:
        app.secret_key = _f.read()
else:
    app.secret_key = os.urandom(32)
    with open(_KEY_FILE, "wb") as _f:
        _f.write(app.secret_key)

# ── الديكوراتور ────────────────────────────────────────────────────────────

def login_required(f):
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated


def admin_required(f):
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))
        if session.get("role") != "admin":
            flash("غير مصرح لك بهذه الصفحة", "danger")
            return redirect(url_for("dashboard"))
        return f(*args, **kwargs)
    return decorated


# ── متغيرات القوالب ────────────────────────────────────────────────────────

@app.context_processor
def inject_globals():
    now_dt = datetime.now()
    alert_count   = 0
    notifications = []

    if "user_id" in __import__("flask").session:
        # ── تنبيهات الغياب ──────────────────────────────
        try:
            absence_alerts = db.get_absent_alerts(now_dt.year, now_dt.month)
            for a in absence_alerts:
                reasons_str = "،  ".join(
                    f"{r[0]}: {r[1]} (الحد: {r[2]})" for r in a.get("reasons", [])
                )
                notifications.append({
                    "type": "danger",
                    "icon": "fa-user-xmark",
                    "title": f"تجاوز الحد — {a.get('full_name', '')}",
                    "msg":   reasons_str or "تجاوز الحد المسموح",
                })
            alert_count = len(absence_alerts)
        except Exception:
            pass

        # ── تنبيه انتهاء الترخيص ────────────────────────
        try:
            import license as _lic
            _key = db.get_setting("license_key", "")
            if _key:
                _res = _lic.verify_license(_key)
                if _res.get("valid"):
                    _days = _res.get("days_left", 99999)
                    if _days < 30:
                        notifications.append({
                            "type": "warning",
                            "icon": "fa-shield-halved",
                            "title": "تجديد الترخيص",
                            "msg":   f"ينتهي الترخيص خلال {_days} يوم",
                        })
                        alert_count += 1
        except Exception:
            pass

    return {
        "today":         date.today().isoformat(),
        "now":           now_dt,
        "alert_count":   alert_count,
        "notifications": notifications,
    }


# ── المصادقة ──────────────────────────────────────────────────────────────

@app.route("/")
def index():
    return redirect(url_for("dashboard") if "user_id" in session else url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if "user_id" in session:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = db.authenticate(username, password)
        if user == "disabled":
            flash("هذا الحساب معطّل. تواصل مع المدير", "warning")
        elif user:
            session.permanent = True
            session["user_id"]   = user["id"]
            session["username"]  = user["username"]
            session["role"]      = user["role"]
            session["full_name"] = user["full_name"]
            if user["role"] == "admin" and db.is_first_run():
                return redirect(url_for("change_password"))
            return redirect(url_for("dashboard"))
        else:
            flash("اسم المستخدم أو كلمة المرور غير صحيحة", "danger")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("تم تسجيل الخروج", "info")
    return redirect(url_for("login"))


# ── لوحة التحكم ────────────────────────────────────────────────────────────

@app.route("/dashboard")
@login_required
def dashboard():
    today  = date.today().isoformat()
    now_dt = datetime.now()
    # سجّل غياب اليوم تلقائياً للموظفين المتجاوزين لحدودهم
    db.auto_fill_overdue_absences(today)
    stats  = db.get_dashboard_stats(today)
    recent = db.get_attendance_for_date(today)
    alerts = db.get_absent_alerts(now_dt.year, now_dt.month)
    emp_stats = db.get_employee_attendance_stats(now_dt.year, now_dt.month)
    return render_template("dashboard.html", stats=stats, recent=recent,
                           today=today, alerts=alerts, emp_stats=emp_stats,
                           current_month=now_dt.strftime('%Y/%m'))


# ── الموظفون ───────────────────────────────────────────────────────────────

@app.route("/employees")
@login_required
def employees():
    show_all = request.args.get("all", "0") == "1"
    emps = db.get_all_employees(include_inactive=show_all)
    return render_template("employees.html", employees=emps, show_all=show_all)


@app.route("/employees/add", methods=["GET", "POST"])
@admin_required
def add_employee():
    if request.method == "POST":
        try:
            db.add_employee(
                request.form["employee_code"].strip(),
                request.form["full_name"].strip(),
                request.form.get("department", "").strip(),
                request.form.get("organization", "").strip(),
                request.form.get("position", "").strip(),
                request.form.get("phone", "").strip(),
                request.form.get("email", "").strip(),
                request.form.get("hire_date", "").strip(),
                int(request.form.get("max_absent_days", 0)),
                int(request.form.get("max_vacation_days", 0)),
                int(request.form.get("max_sick_days", 0)),
                int(request.form.get("max_work_days", 0)),
                int(request.form.get("annual_leave_days", 21)),
                request.form.get("branch", "").strip() or None,
            )
            db.log_action(session["user_id"], session["username"],
                          "إضافة موظف", request.form["full_name"].strip())
            flash("تمت إضافة الموظف بنجاح", "success")
            return redirect(url_for("employees"))
        except Exception as exc:
            flash(f"خطأ: {exc}", "danger")
    branches = db.get_all_branches()
    return render_template("employee_form.html", emp=None, branches=branches)


@app.route("/employees/edit/<int:emp_id>", methods=["GET", "POST"])
@admin_required
def edit_employee(emp_id):
    emp = db.get_employee_by_id(emp_id)
    if not emp:
        flash("الموظف غير موجود", "danger")
        return redirect(url_for("employees"))
    if request.method == "POST":
        try:
            db.update_employee(
                emp_id,
                request.form["employee_code"].strip(),
                request.form["full_name"].strip(),
                request.form.get("department", "").strip(),
                request.form.get("organization", "").strip(),
                request.form.get("position", "").strip(),
                request.form.get("phone", "").strip(),
                request.form.get("email", "").strip(),
                request.form.get("hire_date", "").strip(),
                request.form.get("status", "active"),
                int(request.form.get("max_absent_days", 0)),
                int(request.form.get("max_vacation_days", 0)),
                int(request.form.get("max_sick_days", 0)),
                int(request.form.get("max_work_days", 0)),
                int(request.form.get("annual_leave_days", 21)),
                request.form.get("branch", "").strip() or None,
            )
            db.log_action(session["user_id"], session["username"],
                          "تعديل موظف", request.form["full_name"].strip())
            flash("تم تحديث بيانات الموظف", "success")
            return redirect(url_for("employees"))
        except Exception as exc:
            flash(f"خطأ: {exc}", "danger")
    branches = db.get_all_branches()
    return render_template("employee_form.html", emp=emp, branches=branches)


@app.route("/employees/delete/<int:emp_id>", methods=["POST"])
@admin_required
def delete_employee_route(emp_id):
    emp = db.get_employee_by_id(emp_id)
    name = emp["full_name"] if emp else str(emp_id)
    db.delete_employee(emp_id)
    db.log_action(session["user_id"], session["username"], "حذف موظف", name)
    flash("تم حذف الموظف", "success")
    return redirect(url_for("employees"))


# ── الحضور ────────────────────────────────────────────────────────────────

@app.route("/attendance")
@login_required
def attendance():
    sel_date     = request.args.get("date", date.today().isoformat())
    emps         = db.get_all_employees()
    existing     = {r["employee_id"]: r for r in db.get_attendance_for_date(sel_date)}
    last_statuses = db.get_employees_last_status(sel_date)
    return render_template("attendance.html", employees=emps,
                           existing=existing, sel_date=sel_date,
                           last_statuses=last_statuses)


@app.route("/attendance/save", methods=["POST"])
@login_required
def save_attendance():
    sel_date = request.form.get("date", date.today().isoformat())
    emps     = db.get_all_employees()
    uid      = session["user_id"]
    for emp in emps:
        eid = emp["id"]
        db.record_attendance(
            eid, sel_date,
            request.form.get(f"status_{eid}", "absent"),
            request.form.get(f"check_in_{eid}")  or None,
            request.form.get(f"check_out_{eid}") or None,
            request.form.get(f"exit_type_{eid}") or None,
            request.form.get(f"notes_{eid}")     or None,
            uid,
        )
    flash("تم حفظ سجلات الحضور بنجاح", "success")
    db.log_action(session["user_id"], session["username"],
                  "تسجيل حضور", f"تاريخ: {sel_date}")
    return redirect(url_for("attendance", date=sel_date))


@app.route("/attendance/autofill", methods=["POST"])
@login_required
def autofill_day():
    """ملء يوم واحد تلقائياً بنقل آخر حالة"""
    sel_date = request.form.get("date", date.today().isoformat())
    filled   = db.auto_fill_attendance(sel_date, session["user_id"])
    if filled:
        flash(f"تم التعبئة التلقائية لـ {filled} موظف", "success")
    else:
        flash("جميع الموظفين لديهم سجلات مكتملة لهذا اليوم", "info")
    return redirect(url_for("attendance", date=sel_date))


@app.route("/attendance/fill_range", methods=["POST"])
@login_required
def autofill_range():
    """ملء جميع الأيام المفقودة بين تاريخين"""
    from_date = request.form.get("fill_from", date.today().isoformat())
    to_date   = request.form.get("fill_to",   date.today().isoformat())
    filled    = db.fill_missing_range(from_date, to_date, session["user_id"])
    flash(f"تم التعبئة التلقائية لـ {filled} سجل مفقود بين {from_date} و {to_date}", "success")
    return redirect(url_for("attendance", date=to_date))


@app.route("/attendance/bulk_absent", methods=["POST"])
@login_required
def bulk_absent():
    """تسجيل غياب موظف واحد لعدة أيام متتالية"""
    emp_id    = int(request.form.get("emp_id", 0))
    from_date = request.form.get("from_date", date.today().isoformat())
    num_days  = max(1, min(int(request.form.get("num_days", 1)), 90))
    exit_type = request.form.get("exit_type", "")
    notes     = request.form.get("notes", "").strip()

    from datetime import timedelta
    d0      = date.fromisoformat(from_date)
    saved   = 0
    for i in range(num_days):
        day_str = (d0 + timedelta(days=i)).isoformat()
        db.record_attendance(
            employee_id=emp_id,
            date=day_str,
            status="absent",
            check_in=None,
            check_out=None,
            exit_type=exit_type or None,
            notes=notes or None,
            recorded_by=session["user_id"],
        )
        saved += 1

    flash(f"تم تسجيل الغياب لـ {saved} يوم", "success")

    # تحديث الحد تلقائياً بعدد الأيام المسجّلة
    db.update_absence_limit(emp_id, exit_type or None, saved)

    return redirect(url_for("attendance", date=from_date))



# ── التقارير ──────────────────────────────────────────────────────────────

@app.route("/reports")
@login_required
def reports():
    from_date  = request.args.get("from_date", date.today().isoformat())
    to_date    = request.args.get("to_date",   date.today().isoformat())
    dept       = request.args.get("department", "")
    branch     = request.args.get("branch", "")
    emp_id_str = request.args.get("employee_id", "")
    searched   = "search" in request.args

    data = []
    if searched:
        data = db.get_attendance_report(
            from_date, to_date,
            employee_id=int(emp_id_str) if emp_id_str else None,
            department=dept or None,
            branch=branch or None,
        )

    status_counts = Counter(r["status"] for r in data) if data else {}

    return render_template(
        "reports.html",
        data=data, from_date=from_date, to_date=to_date,
        dept=dept, branch=branch, emp_id=emp_id_str,
        departments=db.get_departments(),
        all_employees=db.get_all_employees(),
        all_branches=db.get_all_branches(),
        searched=searched,
        status_counts=status_counts,
    )


@app.route("/reports/summary")
@login_required
def reports_summary():
    import calendar
    today      = date.today()
    view       = request.args.get("view", "daily")   # daily | weekly | monthly
    # نطاق افتراضي: الشهر الحالي
    default_from = today.replace(day=1).isoformat()
    default_to   = today.isoformat()
    from_date  = request.args.get("from_date", default_from)
    to_date    = request.args.get("to_date",   default_to)

    group_by_map = {"daily": "day", "weekly": "week", "monthly": "month"}
    group_by = group_by_map.get(view, "day")

    rows = db.get_attendance_summary(from_date, to_date, group_by)

    # بيانات الرسم البياني (JSON-safe)
    chart_labels   = [r["period"] for r in rows]
    chart_present  = [r["present"] for r in rows]
    chart_absent   = [r["absent"]  for r in rows]
    chart_late     = [r["late"]    for r in rows]
    chart_leave    = [r["leave"]   for r in rows]
    chart_rate     = [r["rate"]    for r in rows]

    # إجماليات عامة
    totals = {
        "present": sum(r["present"] for r in rows),
        "absent":  sum(r["absent"]  for r in rows),
        "late":    sum(r["late"]    for r in rows),
        "leave":   sum(r["leave"]   for r in rows),
        "total":   sum(r["total"]   for r in rows),
    }
    if totals["total"]:
        totals["rate"] = round(
            (totals["present"] + totals["late"]) / totals["total"] * 100, 1)
    else:
        totals["rate"] = 0

    return render_template(
        "summary.html",
        rows=rows, view=view,
        from_date=from_date, to_date=to_date,
        chart_labels=chart_labels,
        chart_present=chart_present, chart_absent=chart_absent,
        chart_late=chart_late,       chart_leave=chart_leave,
        chart_rate=chart_rate,
        totals=totals,
    )


def _report_data_from_args():
    from_date  = request.args.get("from_date", date.today().isoformat())
    to_date    = request.args.get("to_date",   date.today().isoformat())
    dept       = request.args.get("department", "")
    branch     = request.args.get("branch", "")
    emp_id_str = request.args.get("employee_id", "")
    data = db.get_attendance_report(
        from_date, to_date,
        employee_id=int(emp_id_str) if emp_id_str else None,
        department=dept or None,
        branch=branch or None,
    )
    return from_date, to_date, data


@app.route("/reports/export/excel")
@login_required
def export_excel():
    from_date, to_date, data = _report_data_from_args()
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    eu.export_to_excel(data, from_date, to_date, tmp.name)
    return send_file(
        tmp.name, as_attachment=True,
        download_name=f"attendance_{from_date}_{to_date}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/reports/export/pdf")
@login_required
def export_pdf():
    from_date, to_date, data = _report_data_from_args()
    tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
    tmp.close()
    eu.export_to_pdf(data, from_date, to_date, tmp.name)
    return send_file(
        tmp.name, as_attachment=True,
        download_name=f"attendance_{from_date}_{to_date}.pdf",
        mimetype="application/pdf",
    )


# ── المستخدمون ─────────────────────────────────────────────────────────────

@app.route("/users")
@admin_required
def users():
    return render_template("users.html", users=db.get_all_users())


@app.route("/users/add", methods=["POST"])
@admin_required
def add_user():
    try:
        db.add_user(
            request.form["username"].strip(),
            request.form["password"],
            request.form.get("role", "employee"),
            request.form["full_name"].strip(),
        )
        flash("تمت إضافة المستخدم بنجاح", "success")
    except Exception as exc:
        flash(f"خطأ: {exc}", "danger")
    return redirect(url_for("users"))


@app.route("/users/delete/<int:user_id>", methods=["POST"])
@admin_required
def delete_user(user_id):
    if user_id == session["user_id"]:
        flash("لا يمكنك حذف حسابك الحالي", "danger")
    else:
        db.delete_user(user_id)
        flash("تم حذف المستخدم", "success")
    return redirect(url_for("users"))


@app.route("/users/toggle/<int:user_id>", methods=["POST"])
@admin_required
def toggle_user(user_id):
    if user_id == session["user_id"]:
        flash("لا يمكنك تعطيل حسابك الحالي", "danger")
    else:
        db.toggle_user_active(user_id)
        flash("تم تغيير حالة الحساب", "success")
    return redirect(url_for("users"))


# ── تغيير كلمة المرور ────────────────────────────────────────────────────

@app.route("/change-password", methods=["GET", "POST"])
@login_required
def change_password():
    first_run = db.is_first_run() and session.get("role") == "admin"
    if request.method == "POST":
        current  = request.form.get("current_password", "")
        new_pw   = request.form.get("new_password", "")
        confirm  = request.form.get("confirm_password", "")
        if not first_run:
            user = db.authenticate(session["username"], current)
            if not user:
                flash("كلمة المرور الحالية غير صحيحة", "danger")
                return render_template("change_password.html", first_run=first_run)
        if len(new_pw) < 6:
            flash("كلمة المرور الجديدة يجب أن تكون 6 أحرف على الأقل", "danger")
            return render_template("change_password.html", first_run=first_run)
        if new_pw != confirm:
            flash("كلمة المرور وتأكيدها غير متطابقتين", "danger")
            return render_template("change_password.html", first_run=first_run)
        db.change_password(session["user_id"], new_pw)
        if first_run:
            db.clear_first_run()
        flash("تم تغيير كلمة المرور بنجاح", "success")
        return redirect(url_for("dashboard"))
    return render_template("change_password.html", first_run=first_run)


# ── دليل المستخدم ─────────────────────────────────────────────────────────

@app.route("/help")
@login_required
def help_page():
    return render_template("help.html")


# ── عن البرنامج + نسخ احتياطي ────────────────────────────────────────────

@app.route("/about")
@login_required
def about():
    import license as lic
    backups    = db.list_backups()
    last_bk    = db.get_setting("last_backup", "لم يتم بعد")
    version    = db.get_setting("app_version", "2.0.0")
    lic_status = lic.get_license_status()
    # آخر 30 سطر من ملف الأخطاء
    error_log_lines = []
    try:
        log_path = os.path.join(os.path.dirname(db.DB_PATH), "error.log")
        if os.path.exists(log_path):
            with open(log_path, "r", encoding="utf-8", errors="replace") as f:
                error_log_lines = f.readlines()[-30:]
    except Exception:
        pass
    return render_template("about.html", backups=backups,
                           last_backup=last_bk, version=version,
                           lic_status=lic_status,
                           error_log_lines=error_log_lines)


@app.route("/privacy")
def privacy_policy():
    """صفحة سياسة الخصوصية — مطلوبة لـ Microsoft Store"""
    return render_template("privacy.html")


@app.route("/backup/create", methods=["POST"])
@admin_required
def create_backup():
    try:
        path = db.create_backup()
        flash(f"تم إنشاء النسخة الاحتياطية بنجاح", "success")
    except Exception as e:
        flash(f"فشل إنشاء النسخة الاحتياطية: {e}", "danger")
    return redirect(url_for("about"))


@app.route("/backup/download/<filename>")
@admin_required
def download_backup(filename):
    import re
    if not re.match(r'^attendance_\d{8}_\d{6}\.db$', filename):
        flash("ملف غير صالح", "danger")
        return redirect(url_for("about"))
    backup_dir = os.path.join(os.path.dirname(db.DB_PATH), "backups")
    return send_file(os.path.join(backup_dir, filename),
                     as_attachment=True, download_name=filename)


# ── تقرير الموظف الفردي ─────────────────────────────────────────────────

@app.route("/employees/<int:emp_id>/report")
@login_required
def employee_report(emp_id):
    emp = db.get_employee_by_id(emp_id)
    if not emp:
        flash("الموظف غير موجود", "danger")
        return redirect(url_for("employees"))
    today = date.today()
    year  = int(request.args.get("year",  today.year))
    month = int(request.args.get("month", today.month))
    records, year_stats, emp_full = db.get_employee_full_report(emp_id, year, month)
    month_names = ["","يناير","فبراير","مارس","أبريل","مايو","يونيو",
                   "يوليو","أغسطس","سبتمبر","أكتوبر","نوفمبر","ديسمبر"]
    return render_template("employee_report.html",
        emp=emp_full, records=records, year_stats=year_stats,
        year=year, month=month, month_names=month_names,
        years=list(range(today.year - 3, today.year + 2)))


@app.route("/employees/<int:emp_id>/report/pdf")
@login_required
def employee_report_pdf(emp_id):
    emp = db.get_employee_by_id(emp_id)
    if not emp:
        flash("الموظف غير موجود", "danger")
        return redirect(url_for("employees"))
    today = date.today()
    year  = int(request.args.get("year",  today.year))
    month = int(request.args.get("month", today.month))
    records, year_stats, emp_full = db.get_employee_full_report(emp_id, year, month)
    try:
        tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
        tmp.close()
        eu.export_employee_report_pdf(
            dict(emp_full) if emp_full else {},
            [dict(r) for r in records],
            [dict(s) for s in year_stats],
            year, month, tmp.name
        )
        month_names = ["","يناير","فبراير","مارس","أبريل","مايو","يونيو",
                       "يوليو","أغسطس","سبتمبر","أكتوبر","نوفمبر","ديسمبر"]
        fname = f"تقرير_{emp_full['full_name']}_{year}_{month_names[month]}.pdf"
        return send_file(tmp.name, as_attachment=True,
                         download_name=fname,
                         mimetype="application/pdf")
    except Exception as e:
        flash(f"فشل إنشاء PDF: {e}", "danger")
        return redirect(url_for("employee_report", emp_id=emp_id, year=year, month=month))


@app.route("/employees/<int:emp_id>/id-card")
@login_required
def employee_id_card(emp_id):
    """تحميل بطاقة هوية الموظف بصيغة PDF بحجم بطاقة الائتمان"""
    emp = db.get_employee_by_id(emp_id)
    if not emp:
        flash("الموظف غير موجود", "danger")
        return redirect(url_for("employees"))
    try:
        tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
        tmp.close()
        eu.export_employee_id_card(dict(emp), tmp.name)
        fname = f"بطاقة_{emp['full_name']}.pdf"
        return send_file(tmp.name, as_attachment=True,
                         download_name=fname,
                         mimetype="application/pdf")
    except Exception as e:
        flash(f"فشل إنشاء البطاقة: {e}", "danger")
        return redirect(url_for("employees"))


# ── العطل الرسمية ────────────────────────────────────────────────────────

@app.route("/holidays")
@admin_required
def holidays():
    return render_template("holidays.html", holidays=db.get_all_holidays())


@app.route("/holidays/add", methods=["POST"])
@admin_required
def add_holiday():
    name = request.form.get("name", "").strip()
    d    = request.form.get("date", "").strip()
    rec  = 1 if request.form.get("is_recurring") else 0
    if name and d:
        db.add_holiday(name, d, rec)
        db.log_action(session["user_id"], session["username"],
                      "إضافة عطلة", f"{name} — {d}")
        flash(f"تمت إضافة العطلة: {name}", "success")
    else:
        flash("يجب تعبئة الاسم والتاريخ", "danger")
    return redirect(url_for("holidays"))


@app.route("/holidays/delete/<int:holiday_id>", methods=["POST"])
@admin_required
def delete_holiday(holiday_id):
    db.delete_holiday(holiday_id)
    db.log_action(session["user_id"], session["username"], "حذف عطلة", str(holiday_id))
    flash("تم حذف العطلة", "success")
    return redirect(url_for("holidays"))


# ── الإعدادات ────────────────────────────────────────────────────────────

@app.route("/settings", methods=["GET", "POST"])
@admin_required
def settings():
    if request.method == "POST":
        selected = request.form.getlist("work_days")
        work_days_val = ",".join(selected) if selected else "0,1,2,3,4"
        db.set_setting("work_days", work_days_val)
        db.log_action(session["user_id"], session["username"],
                      "تعديل إعدادات", f"أيام العمل: {work_days_val}")
        flash("تم حفظ الإعدادات بنجاح", "success")
        return redirect(url_for("settings"))
    current_work_days = db.get_work_days()
    return render_template("settings.html",
                           current_work_days={str(d) for d in current_work_days})


# ── استيراد الموظفين من Excel ────────────────────────────────────────────

@app.route("/employees/import", methods=["GET", "POST"])
@admin_required
def import_employees():
    if request.method == "POST":
        f = request.files.get("excel_file")
        if not f or not f.filename.endswith((".xlsx", ".xls")):
            flash("يجب رفع ملف Excel (.xlsx أو .xls)", "danger")
            return redirect(url_for("import_employees"))
        try:
            import openpyxl
            tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
            f.save(tmp.name)
            tmp.close()
            wb = openpyxl.load_workbook(tmp.name, data_only=True)
            ws = wb.active
            headers = [str(c.value or "").strip().lower() for c in ws[1]]
            col_map = {
                "employee_code": ["employee_code","كود الموظف","الكود","code"],
                "full_name":     ["full_name","اسم الموظف","الاسم","الاسم الكامل","name"],
                "department":    ["department","القسم","قسم"],
                "organization":  ["organization","الجهة","جهة"],
                "position":      ["position","المنصب","الوظيفة"],
                "phone":         ["phone","الهاتف","رقم الهاتف"],
                "email":         ["email","البريد","البريد الإلكتروني"],
                "hire_date":     ["hire_date","تاريخ التعيين","تاريخ الالتحاق"],
            }
            idx_map = {}
            for field, aliases in col_map.items():
                for i, h in enumerate(headers):
                    if h in aliases:
                        idx_map[field] = i
                        break
            rows = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not any(row):
                    continue
                r = {}
                for field, idx in idx_map.items():
                    val = row[idx] if idx < len(row) else ""
                    r[field] = str(val) if val is not None else ""
                rows.append(r)
            ok, fail, dup = db.import_employees_from_list(rows)
            db.log_action(session["user_id"], session["username"],
                          "استيراد موظفين", f"نجح:{ok} فشل:{fail} مكرر:{dup}")
            flash(f"تم الاستيراد ← نجح: {ok} | فشل: {fail} | مكرر (موجود): {dup}", "success")
        except ImportError:
            flash("مكتبة openpyxl غير مثبتة. نفّذ: pip install openpyxl", "danger")
        except Exception as exc:
            flash(f"خطأ في قراءة الملف: {exc}", "danger")
        return redirect(url_for("employees"))
    return render_template("import_employees.html")


@app.route("/employees/import/template")
@admin_required
def download_import_template():
    """تحميل نموذج Excel فارغ"""
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "الموظفون"
        headers = ["employee_code","full_name","department","organization",
                   "position","phone","email","hire_date"]
        ws.append(headers)
        example = ["EMP001","أحمد محمد","الموارد البشرية","الإدارة العامة",
                   "مدير","0501234567","ahmed@example.com","2024-01-01"]
        ws.append(example)
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        wb.save(tmp.name)
        tmp.close()
        return send_file(tmp.name, as_attachment=True,
                         download_name="import_template.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except ImportError:
        flash("مكتبة openpyxl غير مثبتة.", "danger")
        return redirect(url_for("import_employees"))


# ── سجل التعديلات ────────────────────────────────────────────────────────

@app.route("/audit-log")
@admin_required
def audit_log():
    logs = db.get_audit_log(300)
    return render_template("audit_log.html", logs=logs)


# ── الفروع ───────────────────────────────────────────────────────────────

@app.route("/branches")
@admin_required
def branches():
    return render_template("branches.html", branches=db.get_all_branches())


@app.route("/branches/add", methods=["POST"])
@admin_required
def add_branch():
    name    = request.form.get("name", "").strip()
    address = request.form.get("address", "").strip()
    manager = request.form.get("manager", "").strip()
    if not name:
        flash("اسم الفرع مطلوب", "danger")
        return redirect(url_for("branches"))
    try:
        db.add_branch(name, address, manager)
        db.log_action(session["user_id"], session["username"], "إضافة فرع", name)
        flash(f"تمت إضافة الفرع: {name}", "success")
    except Exception as e:
        flash(f"خطأ: {e}", "danger")
    return redirect(url_for("branches"))


@app.route("/branches/edit/<int:branch_id>", methods=["POST"])
@admin_required
def edit_branch(branch_id):
    name    = request.form.get("name", "").strip()
    address = request.form.get("address", "").strip()
    manager = request.form.get("manager", "").strip()
    if not name:
        flash("اسم الفرع مطلوب", "danger")
        return redirect(url_for("branches"))
    db.update_branch(branch_id, name, address, manager)
    db.log_action(session["user_id"], session["username"], "تعديل فرع", name)
    flash("تم تحديث الفرع", "success")
    return redirect(url_for("branches"))


@app.route("/branches/delete/<int:branch_id>", methods=["POST"])
@admin_required
def delete_branch(branch_id):
    db.delete_branch(branch_id)
    db.log_action(session["user_id"], session["username"], "حذف فرع", str(branch_id))
    flash("تم حذف الفرع", "success")
    return redirect(url_for("branches"))


# ── الترخيص ──────────────────────────────────────────────────────────────
@app.route("/license", methods=["GET", "POST"])
@admin_required
def license_page():
    import license as lic
    if request.method == "POST":
        key = request.form.get("license_key", "").strip()
        result = lic.save_license(key)
        if result["valid"]:
            db.log_action(session["user_id"], session["username"],
                          "تفعيل ترخيص", f"العميل: {result.get('customer','')}")
            flash(f"✅ تم تفعيل الترخيص بنجاح — {result.get('customer','')}", "success")
        else:
            flash(f"❌ مفتاح غير صالح: {result.get('reason','')}", "danger")
        return redirect(url_for("license_page"))
    status = lic.get_license_status()
    return render_template("license.html", status=status)


# ── معالجات الأخطاء ───────────────────────────────────────────────────────

@app.errorhandler(404)
def page_not_found(e):
    return render_template("404.html"), 404

@app.errorhandler(500)
def internal_server_error(e):
    return render_template("500.html"), 500

@app.errorhandler(403)
def forbidden(e):
    flash("ليس لديك صلاحية للوصول إلى هذه الصفحة", "danger")
    return redirect(url_for("dashboard")), 303


# ── نقطة التشغيل ──────────────────────────────────────────────────────────

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--docker", action="store_true", help="تشغيل في وضع Docker")
    args, _ = parser.parse_known_args()

    db.init_db()

    # نسخ احتياطي تلقائي كل 24 ساعة
    def _auto_backup():
        try:
            db.create_backup()
        except Exception:
            pass
        t = threading.Timer(86400, _auto_backup)
        t.daemon = True
        t.start()

    _backup_timer = threading.Timer(86400, _auto_backup)
    _backup_timer.daemon = True
    _backup_timer.start()

    host = "0.0.0.0"
    port = int(os.environ.get("PORT", 5000))

    if args.docker:
        print(f"  ✔ نظام الحضور والغياب – Docker Mode")
        print(f"  ► افتح المتصفح على: http://localhost:{port}")
    else:
        print(f"\n  ✔ نظام الحضور والغياب – واجهة الويب")
        print(f"  ► افتح المتصفح على: http://localhost:{port}")
        print(f"  ► من الهاتف (نفس الشبكة): http://<IP الحاسوب>:{port}")
        print(f"  ► بيانات الدخول: admin / admin123\n")

    app.run(debug=False, host=host, port=port)
