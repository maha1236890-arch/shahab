"""
أدوات التصدير: PDF و Excel
Export Utilities: PDF and Excel
"""

import os
from PyQt5.QtPrintSupport import QPrinter
from PyQt5.QtGui import QTextDocument
from PyQt5.QtWidgets import QFileDialog, QMessageBox
from datetime import datetime


# ── PDF ─────────────────────────────────────────────────────────────────────

def save_as_pdf(parent, html: str, default_name: str = 'تقرير.pdf',
                page_size=QPrinter.A4) -> bool:
    """حفظ محتوى HTML كملف PDF مع مربع حوار اختيار المسار"""
    path, _ = QFileDialog.getSaveFileName(
        parent, 'حفظ التقرير PDF', default_name, 'PDF (*.pdf)'
    )
    if not path:
        return False
    if not path.lower().endswith('.pdf'):
        path += '.pdf'

    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(path)
    printer.setPageSize(page_size)
    printer.setPageMargins(15, 15, 15, 15, QPrinter.Millimeter)

    doc = QTextDocument()
    doc.setHtml(html)
    doc.print_(printer)

    QMessageBox.information(parent, 'تم الحفظ', f'تم حفظ التقرير بنجاح:\n{path}')
    return True


# ── Excel ────────────────────────────────────────────────────────────────────

def save_as_excel(parent, headers: list, rows: list,
                  default_name: str = 'تقرير.xlsx',
                  sheet_title: str = 'التقرير') -> bool:
    """حفظ بيانات جدولية كملف Excel مع تنسيق احترافي"""
    try:
        import openpyxl
        from openpyxl.styles import (Font, PatternFill, Alignment,
                                     Border, Side, GradientFill)
        from openpyxl.utils import get_column_letter
    except ImportError:
        QMessageBox.warning(parent, 'خطأ',
                            'يجب تثبيت مكتبة openpyxl:\npip install openpyxl')
        return False

    path, _ = QFileDialog.getSaveFileName(
        parent, 'حفظ التقرير Excel', default_name, 'Excel (*.xlsx)'
    )
    if not path:
        return False
    if not path.lower().endswith('.xlsx'):
        path += '.xlsx'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.sheet_view.rightToLeft = True  # RTL

    # ── Styles ──
    hdr_font  = Font(name='Arial', bold=True, color='FFFFFF', size=12)
    hdr_fill  = PatternFill('solid', fgColor='1F6FEB')
    hdr_align = Alignment(horizontal='center', vertical='center',
                          wrap_text=True, readingOrder=2)

    cell_font  = Font(name='Arial', size=11)
    cell_align = Alignment(horizontal='right', vertical='center',
                           wrap_text=True, readingOrder=2)
    alt_fill   = PatternFill('solid', fgColor='EFF6FF')

    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── Title row ──
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1,
                         value=f'{sheet_title}  —  {datetime.now().strftime("%Y/%m/%d  %H:%M")}')
    title_cell.font = Font(name='Arial', bold=True, size=14, color='0D1117')
    title_cell.fill = PatternFill('solid', fgColor='C8D9F5')
    title_cell.alignment = Alignment(horizontal='center', vertical='center',
                                     readingOrder=2)
    ws.row_dimensions[1].height = 28

    # ── Headers row ──
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align
        cell.border = border
    ws.row_dimensions[2].height = 22

    # ── Data rows ──
    for row_idx, row_data in enumerate(rows, 3):
        is_alt = (row_idx % 2 == 0)
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx,
                           value=str(value) if value is not None else '')
            cell.font = cell_font
            cell.alignment = cell_align
            cell.border = border
            if is_alt:
                cell.fill = alt_fill
        ws.row_dimensions[row_idx].height = 18

    # ── Auto column width ──
    for col_idx in range(1, len(headers) + 1):
        max_len = max(
            (len(str(ws.cell(row=r, column=col_idx).value or ''))
             for r in range(1, len(rows) + 3)),
            default=10
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)

    wb.save(path)
    QMessageBox.information(parent, 'تم الحفظ', f'تم حفظ ملف Excel:\n{path}')
    return True


# ── Grouped Excel (for dept-grouped reports) ─────────────────────────────────

def save_grouped_excel(parent, dept_data: dict, date_str: str,
                       default_name: str = 'تقرير_حضور_يومي.xlsx') -> bool:
    """تصدير تقرير الحضور اليومي مقسّمًا حسب الأقسام"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        QMessageBox.warning(parent, 'خطأ', 'يجب تثبيت openpyxl')
        return False

    path, _ = QFileDialog.getSaveFileName(
        parent, 'حفظ تقرير Excel', default_name, 'Excel (*.xlsx)'
    )
    if not path:
        return False
    if not path.lower().endswith('.xlsx'):
        path += '.xlsx'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'الحضور اليومي'
    ws.sheet_view.rightToLeft = True

    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    DEPT_COLORS = [
        '1F6FEB', '238636', '9E6A03', '8250DF',
        'DA3633', '0CA4A5', 'E3B341', 'F78166',
    ]

    current_row = 1

    # ── Main title ──
    ws.merge_cells(start_row=current_row, start_column=1,
                   end_row=current_row, end_column=5)
    tc = ws.cell(row=current_row, column=1,
                 value=f'تقرير الحضور اليومي  —  {date_str}')
    tc.font = Font(name='Arial', bold=True, size=14, color='FFFFFF')
    tc.fill = PatternFill('solid', fgColor='1F6FEB')
    tc.alignment = Alignment(horizontal='center', vertical='center', readingOrder=2)
    ws.row_dimensions[current_row].height = 28
    current_row += 1

    headers = ['الاسم الوظيفي', 'القسم', 'الحالة', 'وقت الحضور', 'الكود']

    for dept_idx, (dept_name, employees) in enumerate(dept_data.items()):
        color_hex = DEPT_COLORS[dept_idx % len(DEPT_COLORS)]

        # ── Department header ──
        ws.merge_cells(start_row=current_row, start_column=1,
                       end_row=current_row, end_column=5)
        dc = ws.cell(row=current_row, column=1,
                     value=f'القسم: {dept_name}  ({len(employees)} موظف)')
        dc.font = Font(name='Arial', bold=True, size=12, color='FFFFFF')
        dc.fill = PatternFill('solid', fgColor=color_hex)
        dc.alignment = Alignment(horizontal='right', vertical='center',
                                 indent=1, readingOrder=2)
        ws.row_dimensions[current_row].height = 22
        current_row += 1

        # ── Column headers ──
        for col_idx, hdr in enumerate(headers, 1):
            c = ws.cell(row=current_row, column=col_idx, value=hdr)
            c.font = Font(name='Arial', bold=True, size=11, color='0D1117')
            c.fill = PatternFill('solid', fgColor='C8D9F5')
            c.alignment = Alignment(horizontal='center', readingOrder=2)
            c.border = border
        ws.row_dimensions[current_row].height = 20
        current_row += 1

        # ── Employee rows ──
        for i, emp in enumerate(employees):
            status = emp.get('status') or 'غير مسجل'
            row_vals = [
                emp.get('job_name', ''),
                dept_name,
                status,
                emp.get('time_in', '') or '',
                emp.get('code', ''),
            ]
            fill_color = ('E8F5E9' if status == 'حاضر'
                          else 'FFEBEE' if status == 'غائب'
                          else 'FAFAFA')
            for col_idx, val in enumerate(row_vals, 1):
                c = ws.cell(row=current_row, column=col_idx, value=val)
                c.font = Font(name='Arial', size=11)
                c.fill = PatternFill('solid', fgColor=fill_color)
                c.alignment = Alignment(horizontal='right', readingOrder=2)
                c.border = border
            ws.row_dimensions[current_row].height = 18
            current_row += 1

        current_row += 1  # blank row between depts

    # ── Column widths ──
    for col_idx, width in enumerate([30, 20, 12, 14, 12], 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(path)
    QMessageBox.information(parent, 'تم الحفظ',
                            f'تم حفظ ملف Excel:\n{path}')
    return True
